#!/usr/bin/env python3
"""Build the Grand Finale deck on the organisers' template: finale/Grenuke_Finale_Deck.pptx.

    python finale/build_deck.py

Our journey in 16 slides (about 8-9 minutes): a title, four section dividers, ten content slides and a close.
One idea per slide, a chart wherever we compare, and the spoken script in the speaker notes (simple English).
Colour system: blue-grey = data and inputs, navy = our models, orange = decisions and the number to remember,
peach = estimates. Animations: a fade between slides; content builds in reading order (fade); diagrams and the
climb chart wipe in from the left. Every number comes from knowledge/numbers.md.
"""
from __future__ import annotations

import shutil
import subprocess
import sys
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_MARKER_STYLE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
TEMPLATE = HERE / "MLC_Presentation_template.pptx"
OUT = HERE / "Grenuke_Finale_Deck.pptx"
ADD_SLIDE = Path.home() / ".claude/skills/synced/d42418a9-4952-416b-af08-af549e2a2647_89c0a8db-9e14-479f-bfbc-9452aba93f36/pptx/scripts/add_slide.py"

NAVY, ORANGE, WHITE = RGBColor(0x16, 0x1D, 0x26), RGBColor(0xFF, 0x62, 0x00), RGBColor(0xFF, 0xFF, 0xFF)
WARM, DATA, PEACH = RGBColor(0xF5, 0xF3, 0xEF), RGBColor(0xD6, 0xDC, 0xE7), RGBColor(0xFF, 0xB2, 0x8B)
MUTED, SOFT, LINE = RGBColor(0x6B, 0x72, 0x80), RGBColor(0xB8, 0xC0, 0xCC), RGBColor(0xE3, 0xE6, 0xEB)
FONT = "Ember Modern Display Standard"
SECTIONS = ["Understanding", "Building", "Feedback", "Results"]


# ---------------------------------------------------------------- drawing helpers

def text(slide, x, y, w, h, paras, size=14, color=NAVY, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         spacing=0):
    """paras: a string, or a list of paragraphs; a paragraph is a string or a list of (text, overrides) runs."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, para in enumerate([paras] if isinstance(paras, str) else paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if spacing and i:
            p.space_before = Pt(spacing)
        for t, o in ([(para, {})] if isinstance(para, str) else para):
            r = p.add_run()
            r.text = t
            r.font.name, r.font.size, r.font.bold = FONT, Pt(o.get("size", size)), o.get("bold", bold)
            r.font.color.rgb = o.get("color", color)
    return tb


def box(slide, x, y, w, h, fill, radius=0.06, shape=MSO_SHAPE.ROUNDED_RECTANGLE, label=None, size=13,
        color=NAVY, bold=True):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = min(radius / min(w, h), 0.5)
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.fill.background()
    s.shadow.inherit = False
    if label is not None:
        tf = s.text_frame
        tf.margin_left = tf.margin_right = Inches(0.05)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = label
        r.font.name, r.font.size, r.font.bold, r.font.color.rgb = FONT, Pt(size), bold, color
    return s


def dot(slide, x, y, d, fill, label, size=12, color=WHITE):
    return box(slide, x, y, d, d, fill, shape=MSO_SHAPE.OVAL, label=label, size=size, color=color)


def chevron(slide, x, y, size=0.16, color=ORANGE):
    return box(slide, x, y, size, size, color, shape=MSO_SHAPE.CHEVRON)


def remove(shape):
    shape._element.getparent().remove(shape._element)


def style_chart(gf, size=11):
    ch = gf.chart
    ch.has_legend = False
    ch.has_title = False
    ch.font.name, ch.font.size = FONT, Pt(size)
    ch.font.color.rgb = NAVY
    return ch


def label_points(series, labels, size=12, color=NAVY, position=XL_LABEL_POSITION.OUTSIDE_END):
    for i, t in enumerate(labels):
        if t is None:
            continue
        dl = series.points[i].data_label
        dl.has_text_frame = True
        dl.text_frame.text = t
        for r in dl.text_frame.paragraphs[0].runs:
            r.font.name, r.font.size, r.font.bold, r.font.color.rgb = FONT, Pt(size), True, color
        dl.position = position


def bars(slide, x, y, w, h, cats, values, colors, vmin, vmax, labels, size=11, gap=70):
    cd = CategoryChartData()
    cd.categories = cats
    cd.add_series("s", values)
    gf = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(x), Inches(y), Inches(w), Inches(h), cd)
    ch = style_chart(gf, size)
    va = ch.value_axis
    va.minimum_scale, va.maximum_scale = vmin, vmax
    va.visible = False
    va.has_major_gridlines = False
    ch.category_axis.format.line.color.rgb = SOFT
    ch.plots[0].gap_width = gap
    ser = ch.plots[0].series[0]
    for i, c in enumerate(colors):
        ser.points[i].format.fill.solid()
        ser.points[i].format.fill.fore_color.rgb = c
    label_points(ser, labels, size=size + 1)
    return gf


def donut(slide, x, y, d, values, colors, hole=68):
    cd = CategoryChartData()
    cd.categories = [str(i) for i in range(len(values))]
    cd.add_series("s", values)
    gf = slide.shapes.add_chart(XL_CHART_TYPE.DOUGHNUT, Inches(x), Inches(y), Inches(d), Inches(d), cd)
    ch = style_chart(gf)
    ser = ch.plots[0].series[0]
    for i, c in enumerate(colors):
        pt = ser.points[i]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = c
        pt.format.line.fill.background()
    plot = ch._chartSpace.find(".//" + qn("c:doughnutChart"))
    hs = plot.find(qn("c:holeSize"))
    if hs is None:
        hs = etree.SubElement(plot, qn("c:holeSize"))
    hs.set("val", str(hole))
    return gf


# ---------------------------------------------------------------- slide frame, notes, animation

def frame(slide, dark, orange, sub, section, principles):
    """The template's two-tone title and sub-header, a section tracker and a quiet leadership-principle tag."""
    for sh in list(slide.shapes):
        if sh.shape_id == 2:
            runs = sh.text_frame.paragraphs[0].runs
            runs[0].text, runs[1].text, runs[2].text, runs[3].text = dark, " ", orange, sub
        elif (sh.has_text_frame and sh.text_frame.text.startswith("Lorem")) or sh.name == "Google Shape;288;p42":
            remove(sh)
    runs = []
    for i, name in enumerate(SECTIONS):
        if i:
            runs.append(("    ", {}))
        cur = name == section
        runs.append((f"0{i + 1} {name}", {"color": ORANGE if cur else SOFT, "bold": cur}))
    text(slide, 0.55, 5.04, 5.5, 0.22, [runs], size=8.5)
    text(slide, 6.0, 5.04, 3.45, 0.22, [[("● ", {"color": ORANGE, "size": 7}), (principles, {})]], size=8.5,
         color=MUTED, align=PP_ALIGN.RIGHT)


def notes(slide, script):
    slide.notes_slide.notes_text_frame.text = script


EFFECTS = {"fade": (10, 0, "fade", 450), "wipe": (22, 8, "wipe(left)", 500)}


def animate(slide, steps):
    """steps: a list of build steps; each step is a list of (shape, 'fade' | 'wipe') that start together.
    The build starts with the slide and each step follows the previous one (After Previous), so no clicks are needed."""
    ns = 'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"'
    nid = [3]

    def new():
        nid[0] += 1
        return nid[0]

    pars, t, built = [], 0, []
    for step in steps:
        effects, longest = [], 0
        for k, (sh, kind) in enumerate(step):
            pid, sub, filt, dur = EFFECTS[kind]
            spid = sh.shape_id
            built.append(sh)
            a, b, c = new(), new(), new()
            node = "afterEffect" if k == 0 else "withEffect"
            effects.append(
                f'<p:par><p:cTn id="{a}" presetID="{pid}" presetClass="entr" presetSubtype="{sub}" fill="hold" grpId="0" '
                f'nodeType="{node}"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
                f'<p:set><p:cBhvr><p:cTn id="{b}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn>'
                f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName>'
                f'</p:attrNameLst></p:cBhvr><p:to><p:strVal val="visible"/></p:to></p:set>'
                f'<p:animEffect transition="in" filter="{filt}"><p:cBhvr><p:cTn id="{c}" dur="{dur}"/><p:tgtEl>'
                f'<p:spTgt spid="{spid}"/></p:tgtEl></p:cBhvr></p:animEffect></p:childTnLst></p:cTn></p:par>')
            longest = max(longest, dur)
        g = new()
        pars.append(f'<p:par><p:cTn id="{g}" fill="hold"><p:stCondLst><p:cond delay="{t}"/></p:stCondLst>'
                    f'<p:childTnLst>{"".join(effects)}</p:childTnLst></p:cTn></p:par>')
        t += longest
    bld = []
    for sh in built:
        if sh.shape_type == 3:      # a chart
            bld.append(f'<p:bldGraphic spid="{sh.shape_id}" grpId="0"><p:bldAsOne/></p:bldGraphic>')
        else:
            bld.append(f'<p:bldP spid="{sh.shape_id}" grpId="0" animBg="1"/>')
    xml = (f'<p:timing {ns}><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>'
           f'<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>'
           f'<p:par><p:cTn id="3" fill="hold"><p:stCondLst><p:cond delay="indefinite"/><p:cond evt="onBegin" delay="0">'
           f'<p:tn val="2"/></p:cond></p:stCondLst><p:childTnLst>{"".join(pars)}</p:childTnLst></p:cTn></p:par>'
           f'</p:childTnLst></p:cTn><p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond>'
           f'</p:prevCondLst><p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond>'
           f'</p:nextCondLst></p:seq></p:childTnLst></p:cTn></p:par></p:tnLst><p:bldLst>{"".join(bld)}</p:bldLst></p:timing>')
    sld = slide._element
    ext = sld.find(qn("p:extLst"))
    timing = etree.fromstring(xml)
    (ext.addprevious(timing) if ext is not None else sld.append(timing))


def transition(slide):
    sld = slide._element
    for old in sld.findall(qn("p:transition")):
        sld.remove(old)
    tr = etree.fromstring('<p:transition xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
                          'spd="med"><p:fade/></p:transition>')
    anchor = sld.find(qn("p:clrMapOvr"))
    (anchor.addnext(tr) if anchor is not None else sld.find(qn("p:cSld")).addnext(tr))


# ---------------------------------------------------------------- slides

def title_slide(s):
    slots = sorted([sh for sh in s.shapes if sh.shape_type == 14], key=lambda sh: sh.left)
    names = sorted([sh for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip() == "NAME HERE"],
                   key=lambda sh: sh.left)
    remove(slots[3]); remove(names[3])
    for i, (slot, name, who) in enumerate(zip(slots[:3], names[:3], ["AMEYA BORKAR", "AARUSH BAKSHI", "SACHI DHOKA"])):
        x = 1.79 + i * 2.24
        slot.left, slot.top, slot.width, slot.height = Inches(x), Inches(2.04), Inches(1.94), Inches(2.21)
        name.left = Inches(x + 0.03)
        name.text_frame.paragraphs[0].runs[0].text = who
    for sh in s.shapes:
        if sh.has_text_frame and sh.text_frame.text.startswith("Team"):
            sh.text_frame.paragraphs[0].runs[1].text = "Grenuke"
    text(s, 0.85, 1.62, 8.3, 0.3, "Business Entity Resolution   ·   Amazon ML Challenge 2026 Grand Finale",
         size=13, color=DATA, align=PP_ALIGN.CENTER)
    notes(s, "(about 15 seconds) Hi, we are team Grenuke: Ameya, Aarush and Sachi. We will walk you through our "
             "journey on this challenge: what the data taught us, what we built, how the leaderboard changed our "
             "thinking, and what we would do next.")


def divider(s, num, title, line, script):
    pill = next(sh for sh in s.shapes if sh.name == "Google Shape;572;p58")
    a = text(s, 1.25, 1.2, 4, 1.1, num, size=72, color=ORANGE, bold=True)
    b = text(s, 1.25, 2.4, 7.6, 0.7, title, size=34, color=WHITE, bold=True)
    c = text(s, 1.25, 3.12, 7.6, 0.4, line, size=15, color=DATA)
    pill.left, pill.top, pill.width, pill.height = Inches(1.25), Inches(3.8), Inches(1.5), Inches(0.34)
    pill.fill.solid()
    pill.fill.fore_color.rgb = ORANGE
    pill.line.fill.background()
    pill.text_frame.text = f"Part {num.lstrip('0')} of 4"
    r = pill.text_frame.paragraphs[0].runs[0]
    r.font.name, r.font.size, r.font.bold, r.font.color.rgb = FONT, Pt(11), True, WHITE
    pill.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    pill.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    animate(s, [[(a, "fade")], [(b, "fade"), (c, "fade"), (pill, "fade")]])
    notes(s, script)


def data_slide(s):
    frame(s, "First, the", "Data", "We read the data before writing any model.", "Understanding",
          "Dive Deep  ·  Learn and Be Curious")
    cd = CategoryChartData()
    cd.categories = ["Training set", "Test set"]
    cd.add_series("Records per entity", (4.68, 5.75))
    cd.add_series("True matches per entity", (3.46, 3.46))
    gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.5), Inches(1.4), Inches(4.9), Inches(2.95), cd)
    ch = style_chart(gf, 11)
    ch.has_legend = True
    ch.legend.position, ch.legend.include_in_layout = XL_LEGEND_POSITION.BOTTOM, False
    ch.legend.font.size = Pt(10)
    va = ch.value_axis
    va.minimum_scale, va.maximum_scale, va.visible, va.has_major_gridlines = 0, 6.6, False, False
    ch.category_axis.format.line.color.rgb = SOFT
    ch.plots[0].gap_width, ch.plots[0].overlap = 60, -10
    for ser, color, labels in zip(ch.plots[0].series, (DATA, NAVY), (["4.68", "5.75"], ["3.46", "≈ 3.46"])):
        ser.format.fill.solid()
        ser.format.fill.fore_color.rgb = color
        label_points(ser, labels, size=12)
    cap = text(s, 0.55, 4.42, 4.85, 0.35, [[("+23% records, same matches  →  ", {}), ("the extras are look-alikes", {"color": ORANGE})]],
               size=12.5, bold=True)
    rows = [("1 in 2", "entities share a name", "the address decides"),
            ("0", "records belong to two entities", "one owner per record"),
            ("15%", "of the test is France, with no labels", "plan for the unknown")]
    built = []
    for i, (num, fact, so) in enumerate(rows):
        y = 1.5 + i * 1.1
        n = text(s, 5.75, y, 1.25, 0.6, num, size=26, color=ORANGE, bold=True)
        f = text(s, 7.05, y + 0.02, 2.4, 0.9, [fact, [("→ " + so, {"bold": True})]], size=12, spacing=3)
        built.append([(n, "fade"), (f, "fade")])
        if i < 2:
            box(s, 5.75, y + 0.98, 3.7, 0.015, LINE, radius=0.005)
    animate(s, [[(gf, "fade")], [(cap, "fade")]] + built)
    notes(s, "(about 60 seconds) When the problem came in, we did not start with a model. We started by reading the "
             "data, because if you don't understand the data, the model is only guessing. This chart shows our first "
             "surprise. The test set has 23 percent more records per entity than training, but about the same number "
             "of true matches. So the extra records are look-alikes, placed there to trick us. We also saw that half of "
             "the entities share their name with another one, so the address has to decide. No record ever belongs to "
             "two entities. And France, 15 percent of the test, had no labels at all.")


def principles_slide(s):
    frame(s, "Two", "Principles", "Two rules we set before building anything.", "Understanding",
          "Customer Obsession  ·  Frugality")
    left = box(s, 0.55, 1.5, 4.3, 3.29, WARM)
    h1 = text(s, 0.8, 1.7, 3.9, 0.4, [[("1  ", {"color": ORANGE}), ("Decide like the metric", {})]], size=17, bold=True)
    sq = [box(s, 1.0, 2.5, 0.62, 0.62, ORANGE, radius=0.08)]
    eq = text(s, 1.72, 2.52, 0.4, 0.55, "=", size=26, bold=True, align=PP_ALIGN.CENTER)
    sq += [box(s, 2.22 + j * 0.62, 2.62, 0.5, 0.5, DATA, radius=0.06) for j in range(4)]
    c1 = text(s, 0.8, 3.35, 3.9, 0.9, ["1 wrong merge costs as much as 4 missed copies.",
                                        [("So: precision first, and an empty answer when unsure.", {"color": MUTED})]],
              size=12.5, spacing=6)
    right = box(s, 5.15, 1.5, 4.3, 3.29, NAVY)
    h2 = text(s, 5.4, 1.7, 3.9, 0.4, [[("2  ", {"color": ORANGE}), ("Compute where it's unsure", {})]], size=17,
              color=WHITE, bold=True)
    d = donut(s, 5.35, 2.2, 1.85, (2.6, 97.4), (ORANGE, RGBColor(0x3A, 0x45, 0x55)))
    pct = text(s, 5.35, 2.88, 1.85, 0.5, "2.6%", size=20, color=ORANGE, bold=True, align=PP_ALIGN.CENTER)
    c2 = text(s, 7.35, 2.6, 1.95, 1.4, ["of pairs ever reach our expensive models.",
                                         [("A cheap model reads every pair first.", {"color": SOFT})]],
              size=12.5, color=WHITE, spacing=6)
    animate(s, [[(left, "fade"), (h1, "fade")], [(x, "fade") for x in sq + [eq]], [(c1, "fade")],
                [(right, "fade"), (h2, "fade")], [(d, "fade"), (pct, "fade")], [(c2, "fade")]])
    notes(s, "(about 45 seconds) From the data, we set two rules. First, decide like the metric. F0.5 punishes one "
             "wrong merge as much as four missed copies, and merging two different businesses is the mistake that hurts "
             "a real customer. So we chose precision first, and an empty answer when we are unsure. Second, spend "
             "compute where the model is unsure. A cheap model reads every pair, and only 2.6 percent of pairs ever "
             "reach our expensive models.")


def blocking_slide(s):
    frame(s, "Smart", "Blocking", "We only compare records that could plausibly match.", "Building",
          "Invent and Simplify")
    layers = [(8.6, DATA, [("Every possible pair  ·  ", {}), ("17 trillion", {"bold": True})], NAVY),
              (5.6, RGBColor(0xC2, 0xCB, 0xD9), [("Retrieval  ·  ", {}), ("58.4M pairs", {"bold": True}),
                                                ("  ·  99.1% of matches kept", {})], NAVY),
              (2.7, NAVY, [("3.7", {"bold": True, "color": ORANGE, "size": 18}), ("  per entity", {"bold": True})], WHITE)]
    built = []
    for i, (w, fill, runs, col) in enumerate(layers):
        x, y = 5.0 - w / 2, 1.5 + i * 0.78
        b = box(s, x, y, w, 0.62, fill, radius=0.1)
        t = text(s, x, y, w, 0.62, [runs], size=13, color=col, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        built.append([(b, "fade"), (t, "fade")])
    side = text(s, 6.5, 3.18, 3.0, 0.5, [[("6.4M pairs  ·  ", {}), ("98.4% kept", {"bold": True})]], size=12)
    built[-1].append((side, "fade"))
    head = text(s, 0.55, 3.95, 3.0, 0.3, "How we search", size=11, color=MUTED, bold=True)
    chips = [("Each country separately", 2.0), ("Name + address words", 1.95), ("Name only, if the address is empty", 2.6),
             ("Repair spellings first", 2.0)]
    x, chip_shapes = 0.55, []
    for label, w in chips:
        chip_shapes.append(box(s, x, 4.25, w - 0.1, 0.42, WARM, radius=0.21, label=label, size=10.5, bold=False))
        x += w + 0.02
    animate(s, built + [[(head, "fade")] + [(c, "fade") for c in chip_shapes]])
    notes(s, "(about 55 seconds) Comparing every entity with every record would mean about 17 trillion pairs. So "
             "first, we search for likely candidates. We search each country separately, using the words in the name "
             "and address, plus a name-only search when the address is empty. Before searching, we repair names: OCR "
             "errors, website-style names, and Hindi spellings. Retrieval keeps 99.1 percent of the true matches. Then "
             "a learned cut leaves just 3.7 candidates per entity, and still keeps 98.4 percent. We chose this soft "
             "matching over hard city or state keys, because many true copies have an empty or landmark-only address.")


def architecture_slide(s):
    frame(s, "Our", "Architecture", "Seven steps, from raw records to final matches.", "Building", "Think Big")
    steps = [("Clean", DATA, NAVY, "fix spellings, OCR digits and web names"),
             ("Retrieve", DATA, NAVY, "search name and address words, per country"),
             ("Score", NAVY, WHITE, "XGBoost on name, address, number and rival features"),
             ("Re-read", NAVY, WHITE, "transformers read the unsure 2.6%"),
             ("Rescore", NAVY, WHITE, "XGBoost adds their votes; calibrated"),
             ("Decide", ORANGE, WHITE, "best set per entity; one owner per record"),
             ("Check", ORANGE, WHITE, "rules for known decoys; a 7B model re-checks")]
    built = []
    for i, (name, fill, col, how) in enumerate(steps):
        x = 0.55 + i * 1.29
        b = box(s, x, 1.75, 1.12, 0.9, fill, radius=0.08, label=name, size=14, color=col)
        h = text(s, x - 0.02, 2.82, 1.16, 1.2, how, size=11, color=NAVY, align=PP_ALIGN.CENTER)
        step = [(b, "wipe"), (h, "fade")]
        if i < len(steps) - 1:
            step.append((chevron(s, x + 1.135, 2.13, size=0.14), "fade"))
        built.append(step)
    legend = []
    for j, (label, fill) in enumerate([("data", DATA), ("models", NAVY), ("decisions", ORANGE)]):
        x = 0.55 + j * 1.5
        legend.append(box(s, x, 4.22, 0.2, 0.2, fill, radius=0.04))
        legend.append(text(s, x + 0.3, 4.19, 1.1, 0.3, label, size=11, color=MUTED))
    note = text(s, 5.2, 4.15, 4.25, 0.5, [[("Cheap steps read everything; ", {}),
                                           ("expensive ones only what they must.", {"bold": True})]],
                size=11.5, align=PP_ALIGN.RIGHT)
    animate(s, built + [[(x, "fade") for x in legend] + [(note, "fade")]])
    notes(s, "(about 60 seconds) Here is our architecture in seven steps. Blue is data, navy is our models, and "
             "orange is where we decide. We clean the records and retrieve candidates. XGBoost, a fast tree model, "
             "scores every pair using name, address, house-number and rival features. The pairs it is unsure about "
             "are re-read by transformer models: e5, bge and Qwen. A second XGBoost combines their votes with context "
             "and gives a calibrated probability. Then we decide per entity. Finally, rules catch known decoys, and a "
             "7B model re-checks the answers we were most sure of.")


def decide_slide(s):
    frame(s, "How a Match Is", "Decided", "One entity, four candidates: a simple example.", "Building",
          "Customer Obsession")
    ent = box(s, 0.55, 1.45, 6.1, 0.62, DATA, radius=0.08)
    et = text(s, 0.75, 1.45, 5.8, 0.62, [[("ENTITY   ", {"size": 9, "color": MUTED}),
                                          ("Acme Robotics Inc  ·  500 Market St, San Jose", {"bold": True})]],
              size=13, anchor=MSO_ANCHOR.MIDDLE)
    cands = [("Acme Robotics Incorporated · 500 Market Street", 0.98, True),
             ("Acme Robotics · Nr. City Hall, San Jose", 0.81, True),
             ("Acme Robotix · 12 Elm Rd", 0.12, False),
             ("Acme Bakery · 500 Market St", 0.03, False)]
    built = [[(ent, "fade"), (et, "fade")]]
    for i, (name, p, keep) in enumerate(cands):
        y = 2.3 + i * 0.6
        n = text(s, 0.55, y, 3.1, 0.45, name, size=11, anchor=MSO_ANCHOR.MIDDLE)
        track = box(s, 3.75, y + 0.12, 2.2, 0.22, LINE, radius=0.11)
        fill = box(s, 3.75, y + 0.12, max(2.2 * p, 0.12), 0.22, NAVY if keep else SOFT, radius=0.11)
        v = text(s, 6.02, y, 0.45, 0.45, f"{p:.2f}", size=11, bold=True, anchor=MSO_ANCHOR.MIDDLE)
        m = text(s, 6.42, y, 0.3, 0.45, "✓" if keep else "✗", size=15, color=ORANGE if keep else MUTED, bold=True,
                 anchor=MSO_ANCHOR.MIDDLE)
        built.append([(n, "fade"), (track, "fade"), (fill, "wipe"), (v, "fade")])
        built[-1].append((m, "fade"))
    cap = text(s, 0.55, 4.72, 6.1, 0.25, "Illustrative probabilities", size=9, color=MUTED)
    panel = box(s, 7.0, 1.45, 2.45, 3.34, NAVY)
    rules = [("1", "Score each pair: a calibrated probability"),
             ("2", "Add matches while the expected F0.5 goes up"),
             ("3", "Each record goes to one entity only")]
    pr = [(panel, "fade")]
    for i, (k, t) in enumerate(rules):
        y = 1.7 + i * 1.0
        pr.append((dot(s, 7.2, y, 0.36, ORANGE, k, size=12), "fade"))
        pr.append((text(s, 7.7, y - 0.02, 1.6, 0.9, t, size=11.5, color=WHITE), "fade"))
    animate(s, built + [[(cap, "fade")], pr])
    notes(s, "(about 55 seconds) Here is how one decision works, with a simple example. Each candidate gets a "
             "calibrated probability. We sort the candidates and keep adding matches while the expected F0.5 for this "
             "entity goes up. The two real copies are kept. The look-alike with a similar name, and the bakery at the "
             "same address, are dropped. And every record can belong to only one entity, so two entities can never "
             "claim the same record. These probabilities are illustrative.")


def leaderboard_slide(s):
    frame(s, "The Leaderboard", "Talked Back", "Our first scores showed a gap our validation could not see.", "Feedback",
          "Insist on the Highest Standards")
    t1 = text(s, 0.55, 1.45, 4.2, 0.3, "Our first scores", size=12, color=MUTED, bold=True)
    c1 = bars(s, 0.45, 1.7, 4.3, 2.55, ["Our validation", "Leaderboard"], (0.9888, 0.9796), (DATA, ORANGE),
              0.965, 0.993, ["0.9888", "0.9796"], size=11)
    t2 = text(s, 5.15, 1.45, 4.3, 0.3, "By country (estimated)", size=12, color=MUTED, bold=True)
    c2 = bars(s, 5.05, 1.7, 4.4, 2.55, ["US", "India", "France"], (0.99, 0.99, 0.93), (NAVY, NAVY, PEACH),
              0.88, 1.005, ["≈ 0.99", "≈ 0.99", "≈ 0.93"], size=11)
    line = text(s, 0.55, 4.42, 8.9, 0.35, [[("We ruled out bugs, split the gap by country, then ", {}),
                                            ("read the French errors by hand.", {"bold": True})]], size=12.5)
    animate(s, [[(t1, "fade"), (c1, "fade")], [(t2, "fade"), (c2, "fade")], [(line, "fade")]])
    notes(s, "(about 55 seconds) Our first uploads were humbling. Our validation said 0.989, but the leaderboard said "
             "0.980. We treated that gap like a bug report. First, we ruled out bugs in the metric and our files. Then "
             "we split the gap by country. The US and India matched our validation, so the whole gap was France, at "
             "about 0.93. Then we read the French mistakes by hand, and found the decoys our model was accepting: one "
             "word of the name swapped, at the same address.")


def france_slide(s):
    frame(s, "A Country We", "Never Saw", "France had no labels, so we let the data teach us, with guards.", "Feedback",
          "Are Right, A Lot")
    steps = [("Learn its words", "word statistics, without labels"),
             ("Teach itself", "learn from its own confident answers, cross-checked"),
             ("Second opinion", "different models vote; a 7B re-checks")]
    built = []
    for i, (h, b) in enumerate(steps):
        y = 1.55 + i * 0.95
        built.append([(dot(s, 0.55, y, 0.46, ORANGE, str(i + 1), size=14), "fade"),
                      (text(s, 1.2, y - 0.02, 4.4, 0.85, [[(h, {"bold": True, "size": 15})], b], size=12), "fade")])
    honest = text(s, 0.55, 4.45, 5.2, 0.3, [[("Honestly: ", {"bold": True}), ("self-training failed twice before it worked.", {})]],
                  size=12, color=MUTED)
    t = text(s, 6.15, 1.45, 3.3, 0.3, "France (estimated)", size=12, color=MUTED, bold=True)
    c = bars(s, 6.05, 1.7, 3.4, 2.65, ["Before", "After"], (0.93, 0.98), (PEACH, ORANGE), 0.88, 1.0,
             ["≈ 0.93", "≈ 0.98"], size=11, gap=60)
    cap = text(s, 6.15, 4.42, 3.3, 0.35, "never measured directly", size=10, color=MUTED, align=PP_ALIGN.CENTER)
    animate(s, built + [[(honest, "fade")], [(t, "fade"), (c, "fade"), (cap, "fade")]])
    notes(s, "(about 60 seconds) France had no labels, so we let the data teach us, carefully. First, we learned which "
             "French words signal a look-alike, without any labels. Second, the model learned from its own most "
             "confident French answers, with checks so that no pair ever grades itself. Honestly, this failed twice "
             "before it worked. Third, we asked for a second opinion: different model families vote, and a 7B model "
             "re-checks confident answers. We estimate France rose from about 0.93 to about 0.98. It is an estimate, "
             "because France was never measured directly.")


def team_slide(s):
    frame(s, "Every Idea", "Tested", "Three people, 329 experiments, and only evidence decided.", "Results",
          "Bias for Action  ·  Ownership  ·  Earn Trust")
    d = donut(s, 0.55, 1.45, 2.75, (177, 126, 26), (ORANGE, NAVY, DATA), hole=66)
    centre = text(s, 0.55, 2.42, 2.75, 0.8, [[("329", {"size": 26, "bold": True})], [("experiments", {"size": 11, "color": MUTED})]],
                  align=PP_ALIGN.CENTER)
    key = text(s, 0.55, 4.35, 3.2, 0.3, [[("■ ", {"color": ORANGE}), ("177 kept   ", {}), ("■ ", {"color": NAVY}),
                                          ("126 dropped   ", {}), ("■ ", {"color": DATA}), ("26 other", {})]], size=10.5)
    team = [("Ameya", "data, blocking, the pipeline, France"),
            ("Aarush", "the 7B model, the re-check, the final submission"),
            ("Sachi", "testing gates, a Qwen model, synthetic French")]
    built = [[(d, "fade"), (centre, "fade")], [(key, "fade")]]
    for i, (n, r) in enumerate(team):
        y = 1.55 + i * 0.78
        built.append([(text(s, 4.0, y, 5.4, 0.7, [[(n, {"color": ORANGE, "bold": True, "size": 16})], r], size=12.5), "fade")])
    rule = box(s, 4.0, 3.95, 5.45, 0.55, WARM, radius=0.27, label="Our rule: when two ideas tie, keep the simpler one.",
               size=12, bold=False)
    animate(s, built + [[(rule, "fade")]])
    notes(s, "(about 50 seconds) We were three people, so we tried a lot. We ran 329 measured experiments. We kept 177 "
             "and dropped 126. Every leaderboard upload tested one change. Ameya led the data work, blocking and France. "
             "Aarush trained the 7B model, built the re-check and the final submission. Sachi built our testing gates, "
             "a Qwen model and synthetic French data. And one simple rule kept us honest: when two ideas tie, keep the "
             "simpler one.")


def climb_slide(s):
    frame(s, "The", "Climb", "Each upload changed one thing, and we learned from each.", "Results", "Deliver Results")
    cats = ["first\nmodel", "+ legal\nforms", "+ France\nrules", "+ stage 3,\nrepairs", "+ trans-\nformers",
            "+ self-\ntraining", "+ decision\nlayer", "+ French\nmodels", "+ France\nlayer", "+ 7B"]
    vals = [0.97608, 0.97961, 0.98781, 0.988609, 0.989721, 0.990179, 0.990264, 0.990545, 0.990699, 0.990879]
    cd = CategoryChartData()
    cd.categories = cats
    cd.add_series("public leaderboard", vals)
    gf = s.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(0.45), Inches(1.42), Inches(6.35), Inches(3.5), cd)
    ch = style_chart(gf, 8.5)
    ch.font.color.rgb = MUTED
    va = ch.value_axis
    va.minimum_scale, va.maximum_scale, va.major_unit = 0.975, 0.992, 0.005
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = LINE
    va.tick_labels.number_format, va.tick_labels.number_format_is_linked = "0.000", False
    va.format.line.fill.background()
    ch.category_axis.format.line.color.rgb = SOFT
    ser = ch.plots[0].series[0]
    ser.smooth = False
    ser.format.line.color.rgb = ORANGE
    ser.format.line.width = Pt(2.75)
    ser.marker.style, ser.marker.size = XL_MARKER_STYLE.CIRCLE, 7
    ser.marker.format.fill.solid()
    ser.marker.format.fill.fore_color.rgb = NAVY
    ser.marker.format.line.fill.background()
    label_points(ser, ["0.976", None, "0.988", None, "0.990", None, None, None, None, "0.991"], size=10,
                 position=XL_LABEL_POSITION.ABOVE)
    card = box(s, 7.0, 1.5, 2.45, 3.29, NAVY)
    c1 = text(s, 7.22, 1.75, 2.1, 0.25, "FINAL SCORE", size=9, color=PEACH, bold=True)
    c2 = text(s, 7.22, 2.0, 2.1, 0.55, "0.990879", size=25, color=ORANGE, bold=True)
    c3 = text(s, 7.22, 2.52, 2.1, 0.3, "public leaderboard", size=11.5, color=SOFT)
    c4 = text(s, 7.22, 3.2, 2.1, 0.7, [[("2nd", {"size": 34, "bold": True}), ("  of the Top 10", {"size": 13})]],
              color=WHITE)
    c5 = text(s, 7.22, 3.95, 2.1, 0.3, "among 32,000+ teams", size=11.5, color=PEACH)
    animate(s, [[(gf, "wipe")], [(card, "fade"), (c1, "fade"), (c2, "fade"), (c3, "fade")], [(c4, "fade"), (c5, "fade")]])
    notes(s, "(about 45 seconds) This is our climb on the public leaderboard, from 0.976 to 0.991. Each point is one "
             "upload, and each one changed one thing we understood. The biggest jump came from the France fixes. "
             "Transformers, self-training and the 7B model each added a smaller step. Our final score was 0.990879, and "
             "we finished second of the top ten, among more than 32,000 teams.")


def lessons_slide(s):
    frame(s, "What We", "Learned", "And what we would do next, at Amazon scale.", "Results",
          "Success and Scale Bring Broad Responsibility")
    left = box(s, 0.55, 1.5, 4.3, 3.29, WARM)
    lh = text(s, 0.8, 1.7, 3.9, 0.35, "Lessons", size=15, color=ORANGE, bold=True)
    right = box(s, 5.15, 1.5, 4.3, 3.29, NAVY)
    rh = text(s, 5.4, 1.7, 3.9, 0.35, "Next, at scale", size=15, color=ORANGE, bold=True)
    lessons = [("Data first", "it drove every good decision"), ("Doubt estimates", "they share the model's blind spots"),
               ("Diversity wins", "most of all in the unknown")]
    nexts = [("Label a little", "500 French pairs where models disagree"), ("Shard the work", "by country and region"),
             ("Guard the ratio", "keep 3 to 4 candidates per entity")]
    built = [[(left, "fade"), (lh, "fade")]]
    for i, (a, b) in enumerate(lessons):
        built.append([(text(s, 0.8, 2.2 + i * 0.82, 3.9, 0.75, [[(a, {"bold": True, "size": 15})], b], size=12), "fade")])
    built.append([(right, "fade"), (rh, "fade")])
    for i, (a, b) in enumerate(nexts):
        built.append([(text(s, 5.4, 2.2 + i * 0.82, 3.9, 0.75, [[(a, {"bold": True, "size": 15})], b], size=12,
                            color=WHITE), "fade")])
    animate(s, built)
    notes(s, "(about 50 seconds) Three lessons. Read the data first: it drove every good decision we made. Doubt your "
             "own estimates: ours were built on our own model, so they shared its blind spots, and the leaderboard "
             "corrected us. And diversity wins: different models together beat one strong model, especially in a "
             "country you have never seen. At Amazon scale, we would label a few hundred French pairs where our models "
             "disagree, split the work by country and region, and keep the candidates per entity small, because every "
             "later step pays for each pair.")


def closing_slide(s):
    pill = next(sh for sh in s.shapes if sh.name == "Google Shape;572;p58")
    a = text(s, 0.53, 1.15, 8.95, 0.9, [[("Thank ", {}), ("You", {"color": ORANGE})]], size=44, color=WHITE, bold=True,
             align=PP_ALIGN.CENTER)
    b = text(s, 1.0, 2.35, 8.0, 0.4, "Read the data.   Decide like the metric.   Spend compute where it matters.",
             size=15, color=DATA, align=PP_ALIGN.CENTER)
    pill.fill.solid()
    pill.fill.fore_color.rgb = ORANGE
    pill.line.fill.background()
    pill.text_frame.text = "Questions welcome"
    r = pill.text_frame.paragraphs[0].runs[0]
    r.font.name, r.font.size, r.font.bold, r.font.color.rgb = FONT, Pt(12), True, WHITE
    pill.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    pill.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    c = text(s, 0.53, 4.1, 8.95, 0.3, "Team Grenuke  ·  Ameya Borkar  ·  Aarush Bakshi  ·  Sachi Dhoka", size=11,
             color=SOFT, align=PP_ALIGN.CENTER)
    animate(s, [[(a, "fade")], [(b, "fade")], [(pill, "fade"), (c, "fade")]])
    notes(s, "(about 10 seconds) To sum up: read the data, decide the way the metric scores, and spend compute where "
             "it matters. Thank you. We are happy to take your questions.")


DIVIDERS = [("01", "Understanding the problem", "What the data told us, and the rules we set ourselves.",
             "(about 5 seconds) Let's start with how we understood the problem."),
            ("02", "Building the solution", "How we search, score and decide.",
             "(about 5 seconds) Now, how we built the solution."),
            ("03", "Learning from feedback", "What the leaderboard showed us, and how we responded.",
             "(about 5 seconds) Then the leaderboard started teaching us."),
            ("04", "Results and reflection", "How we worked, where we landed, and what we learned.",
             "(about 5 seconds) Finally, how we worked and where we landed.")]


def main() -> int:
    work = OUT.with_suffix(".tmp.pptx")
    shutil.copy(TEMPLATE, work)
    for src, n in (("slide2.xml", 9), ("slide3.xml", 4)):      # 10 white content slides, 5 dark slides
        for _ in range(n):
            subprocess.run([sys.executable, str(ADD_SLIDE), str(work), src, "--after", src], check=True, capture_output=True)
    prs = Presentation(work)
    lst = prs.slides._sldIdLst
    ids = list(lst)                                           # 0 title, 1-10 white, 11-15 dark
    for el in ids:
        lst.remove(el)
    for i in [0, 11, 1, 2, 12, 3, 4, 5, 13, 6, 7, 14, 8, 9, 10, 15]:
        lst.append(ids[i])
    sl = list(prs.slides)
    plan = [title_slide, DIVIDERS[0], data_slide, principles_slide, DIVIDERS[1], blocking_slide, architecture_slide,
            decide_slide, DIVIDERS[2], leaderboard_slide, france_slide, DIVIDERS[3], team_slide, climb_slide,
            lessons_slide, closing_slide]
    assert len(sl) == len(plan), (len(sl), len(plan))
    for slide, step in zip(sl, plan):
        divider(slide, *step) if isinstance(step, tuple) else step(slide)
        transition(slide)
    prs.save(OUT)
    work.unlink()
    print(f"{OUT.name}: {len(sl)} slides")
    return 0


if __name__ == "__main__":
    sys.exit(main())
